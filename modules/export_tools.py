def batch_render(core, timelines, preset_name, target_dir):
    from pathlib import Path
    if not core.project: return False, "Connect to a Resolve project first."
    if not target_dir.strip(): return False, "Choose an output directory."
    if not timelines: return False, "Select at least one timeline."
    original = core.project.GetCurrentTimeline()
    added = []
    try:
        Path(target_dir).expanduser().mkdir(parents=True, exist_ok=True)
        available = {t.GetName(): t for t in [core.project.GetTimelineByIndex(i) for i in range(1, core.project.GetTimelineCount()+1)] if t}
        if any(name not in available for name in timelines): return False, "A selected timeline no longer exists. Refresh and select again."
        if preset_name.strip() and not core.project.LoadRenderPreset(preset_name.strip()):
            return False, f"Render preset '{preset_name}' was not found. Import the bundled XML in Resolve, or choose an existing preset."
        for name in dict.fromkeys(timelines):
            timeline = available[name]
            if timeline.GetEndFrame() <= timeline.GetStartFrame(): raise ValueError(f"Timeline is empty: {name}")
            if not core.project.SetCurrentTimeline(timeline): raise ValueError(f"Could not select {name}")
            if not core.project.SetRenderSettings({"TargetDir": str(Path(target_dir).expanduser().resolve()), "CustomName": name, "SelectAllFrames": True}):
                raise ValueError('Resolve rejected the render settings')
            job = core.project.AddRenderJob()
            if not job: raise ValueError(f"Could not queue {name}")
            added.append(job)
        return True, f"Added {len(added)} jobs to the Render Queue. Open Deliver to start rendering."
    except Exception as exc:
        for job in added: core.project.DeleteRenderJob(job)
        return False, str(exc)
    finally:
        if original: core.project.SetCurrentTimeline(original)


def extract_thumbnails(core, mode, target_dir):
    from pathlib import Path
    if not core.project: return False, "Connect to a Resolve project first."
    timeline = core.project.GetCurrentTimeline()
    if not timeline: return False, "No timeline open."
    if not target_dir.strip(): return False, "Choose an output directory."
    if mode not in ('Markers', 'Timeline Center'): return False, "Unknown thumbnail mode."
    original = timeline.GetCurrentTimecode()
    original_page = core.resolve.GetCurrentPage()
    try:
        if not core.resolve.OpenPage('color'):
            return False, 'Open the Color page in Resolve before exporting stills.'
        Path(target_dir).expanduser().mkdir(parents=True, exist_ok=True)
        origin, end = int(timeline.GetStartFrame()), int(timeline.GetEndFrame())
        if end <= origin: return False, "Timeline is empty."
        frames = [origin+int(frame) for frame in sorted(timeline.GetMarkers() or {})] if mode == 'Markers' else [(origin+end)//2]
        frames = [frame for frame in frames if origin <= frame < end]
        if not frames: return False, "No markers inside the timeline."
        fps = float(timeline.GetSetting('timelineFrameRate') or 30)
        safe_name = ''.join(c if c.isalnum() or c in ' _-' else '_' for c in timeline.GetName())
        from time_parser import frames_to_timecode
        for index, frame in enumerate(frames, 1):
            tc = frames_to_timecode(frame, fps, ';' in original)
            if not timeline.SetCurrentTimecode(tc): return False, f"Could not seek to {tc}."
            path = Path(target_dir).expanduser() / f'{safe_name}_Thumbnail-{index:02d}.jpg'
            if not core.project.ExportCurrentFrameAsStill(str(path.resolve())):
                return False, f"Failed to export frame {index}. {index-1} frames saved."
        return True, f"Exported {len(frames)} still frames to {target_dir}."
    except Exception as exc:
        return False, str(exc)
    finally:
        if original: timeline.SetCurrentTimecode(original)
        if original_page: core.resolve.OpenPage(original_page)


def generate_youtube_chapters(core):
    try:
        timeline = core.project.GetCurrentTimeline()
        if not timeline: return False, "No timeline open."
        
        markers = timeline.GetMarkers()
        if not markers: return False, "No markers found."
        
        fps = float(timeline.GetSetting("timelineFrameRate") or 30.0)
        start_frame = int(timeline.GetStartFrame())
        
        chapters = []
        for frame_id, marker in sorted(markers.items(), key=lambda x: int(x[0])):
            rel_frame = int(frame_id)
            rel_sec = max(0, rel_frame / fps)
            
            h = int(rel_sec // 3600)
            m = int((rel_sec % 3600) // 60)
            s = int(rel_sec % 60)
            
            tc_str = f"{h}:{m:02d}:{s:02d}" if h > 0 else f"{m}:{s:02d}"
            # Format: '0:00 MarkerName' with a space after timecode. If no name, just '0:00 '
            marker_name = marker['name'] if marker['name'] else ''
            chapters.append(f"{tc_str} {marker_name}")
            
        if not any(c.startswith("0:00") or c.startswith("00:00") for c in chapters):
            chapters.insert(0, "0:00 Intro")
            
        return True, "\n".join(chapters)
    except Exception as e:
        return False, f"Error: {str(e)}"

def export_csv_data(core, output_path):
    import csv
    try:
        timeline = core.project.GetCurrentTimeline()
        if not timeline: return False, "No timeline open."
        
        items = timeline.GetItemListInTrack("video", 1)
        if not items: return False, "No video clips found."
        
        with open(output_path, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            writer.writerow(["Index", "Clip Name", "Source Start TC", "Duration (Frames)", "Color", "Flags"])
            
            for i, item in enumerate(items, 1):
                name = item.GetName()
                duration = item.GetDuration()
                color = item.GetClipColor()
                flags = ", ".join(item.GetFlagList()) if item.GetFlagList() else ""
                
                mpi = item.GetMediaPoolItem()
                source_tc = mpi.GetClipProperty("Start TC") if mpi else ""
                
                writer.writerow([i, name, source_tc, duration, color, flags])
                
        return True, f"✓ Exported {len(items)} clips data to:\n{output_path}"
    except Exception as e:
        return False, f"Error: {str(e)}"

def frames_to_tc(frames, fps=30):
    from time_parser import frames_to_timecode
    return frames_to_timecode(frames, fps)

def export_shotlist_doc(core, format, output_path, template_path=""):
    import csv
    import os
    try:
        timeline = core.project.GetCurrentTimeline()
        if not timeline: return False, "No timeline open."
        
        if not output_path.strip(): return False, "Choose a shotlist output directory or file."
        if format not in ('csv', 'docx'): return False, "Choose CSV or DOCX format."
        output_path = os.path.expanduser(output_path)
        ext = os.path.splitext(output_path)[1]
        if os.path.isdir(output_path) or not ext:
            if not os.path.exists(output_path):
                os.makedirs(output_path, exist_ok=True)
            project_name = core.project.GetName()
            safe_project_name = "".join([c for c in project_name if c.isalpha() or c.isdigit() or c==' ']).rstrip()
            output_path = os.path.join(output_path, f"{safe_project_name} Shotlist.{format}").replace("\\", "/")
        
        items = timeline.GetItemListInTrack("video", 1)
        if not items: return False, "No video clips found on Track 1."
        
        fps = float(timeline.GetSetting("timelineFrameRate") or 30.0)
        
        shots = []
        for i, item in enumerate(items, 1):
            name = item.GetName()
            
            # Try to guess shot meaning
            meaning = ""
            name_lower = name.lower()
            if "wide" in name_lower or "ws" in name_lower: meaning = "Wide"
            elif "mid" in name_lower or "ms" in name_lower: meaning = "Mid"
            elif "close" in name_lower or "cu" in name_lower: meaning = "Close"
            
            start_frame = int(item.GetStart())
            end_frame = int(item.GetEnd())
            
            start_tc = frames_to_tc(start_frame, fps)
            end_tc = frames_to_tc(end_frame, fps)
            duration_frames = int(item.GetDuration())
            duration_tc = frames_to_tc(duration_frames, fps)
            
            shots.append({
                "serial": i,
                "type": meaning,
                "start_tc": start_tc,
                "end_tc": end_tc,
                "duration": duration_tc,
                "desc": name
            })
            
        if format == 'docx':
            try:
                from docxtpl import DocxTemplate
            except ImportError:
                return False, "docxtpl library is not installed. Run 'pip install docxtpl python-docx'."
                
            if not template_path:
                from docx import Document
                doc = Document()
                doc.add_heading(core.project.GetName() + ' — Shotlist', 0)
                doc.add_paragraph(timeline.GetName())
                table = doc.add_table(rows=1, cols=5)
                for cell, title in zip(table.rows[0].cells, ['ID', 'Clip', 'Start', 'End', 'Duration']): cell.text = title
                for shot in shots:
                    for cell, value in zip(table.add_row().cells, [shot['serial'], shot['desc'], shot['start_tc'], shot['end_tc'], shot['duration']]): cell.text = str(value)
                doc.save(output_path)
                return True, f'Exported Word shotlist to {output_path}'
            if not os.path.isfile(template_path):
                return False, f"Template not found at: {template_path}"
            doc = DocxTemplate(template_path)
            context = {
                "project_name": core.project.GetName(),
                "timeline_name": timeline.GetName(),
                "shots": shots
            }
            doc.render(context)
            doc.save(output_path)
            return True, f"✓ Generated Word Shotlist to:\n{output_path}"
            
        else: # Default CSV
            with open(output_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                writer.writerow(["ID", "Slate/Log Number", "Timecode (Start - End)", "Duration", "Shot Type", "Description", "Primary Tag"])
                for s in shots:
                    tc_range = f"{s['start_tc']} - {s['end_tc']}"
                    writer.writerow([s["serial"], "", tc_range, s["duration"], s["type"], s["desc"], ""])
            return True, f"✓ Exported CSV Shotlist for {len(items)} clips to:\n{output_path}"
            
    except Exception as e:
        return False, f"Error exporting Shotlist: {str(e)}"
